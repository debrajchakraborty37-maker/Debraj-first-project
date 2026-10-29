#!/usr/bin/env python3
"""Daewoo x ShootOrder daily deep-analysis agent (engine).

Input : the finished daily "Campaign Performance Report" workbook (the file you build today).
Output: <outdir>/Deep_Analysis_Report-<end date>.xlsx, digest.md, facts.json
        and an appended snapshot in history/daily_snapshot.csv (enables true daily deltas).

Usage : python3 pipeline/analyze_report.py <report.xlsx> [--outdir outputs/<date>] [--history history/daily_snapshot.csv]
"""
import argparse, json, os, subprocess, sys, shutil
import numpy as np, pandas as pd, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule

NAVY = "1F3A5F"; GREY = "F2F4F7"
HF = PatternFill("solid", fgColor=NAVY); HFont = Font(bold=True, color="FFFFFF", name="Arial", size=10)
TOT = PatternFill("solid", fgColor="E8ECF1")
base = Font(name="Arial", size=10); bold = Font(name="Arial", size=10, bold=True)
inp = Font(name="Arial", size=10, color="0000FF", bold=True)
thin = Side(style="thin", color="D0D5DD"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
INR = '"₹"#,##0'; PCT = '0.0%'; INT = '#,##0'; FRQ = '0.00'
GROUPS = ['Qualified', 'Open / not yet worked', 'Unreachable', 'Junk', 'Not interested / Lost']
DEFAULT_TARGETS = {"cost_per_meeting": 12000, "cpqk": 2500, "qualified_pct": 0.30, "qual_to_meeting": 0.25,
                   "cpl": 700, "contact_rate": 0.70, "junk_rate": 0.05, "ctr": 0.007, "click_to_lead": 0.06,
                   "frequency": 3, "speed_to_lead_min": 15}
inr = lambda x: f"₹{x:,.0f}"
pc = lambda x: f"{x*100:.1f}%"


# ----------------------------------------------------------------------------- load
def num(x): return x if isinstance(x, (int, float)) else 0
def fnum(x):
    try: return float(x)
    except Exception: return x

def load(path):
    wbv = openpyxl.load_workbook(path, data_only=True)
    ws = wbv["Campaign Report"]; rows = list(ws.iter_rows(values_only=True))
    h = [k for k, r in enumerate(rows) if r[0] == 'Reporting starts'][0]
    camp = []
    for r in rows[h + 1:]:
        if r[0] and str(r[0]).startswith('2'):
            camp.append(dict(camp=r[2], status=r[3], spend=num(r[9]), impr=num(r[10]), reach=num(r[11]), leads=num(r[14]),
                             q=num(r[15]), meet=num(r[17]), clicks=num(r[19]), budget=fnum(r[7])))
        elif camp: break
    camp = pd.DataFrame(camp)
    h = [k for k, r in enumerate(rows) if r[0] == 'Parent Campaign' and r[1] == 'Ad set name'][0]
    out = []
    for r in rows[h + 1:]:
        if r[0] is None or str(r[0]).startswith(('Note', 'Subtotal')): continue
        if str(r[0]).startswith('GRAND'): break
        out.append(dict(camp=r[0], adset=r[1], status=r[2], spend=num(r[8]), impr=num(r[9]), reach=num(r[10]),
                        meet=num(r[17]), clicks=num(r[19]), leads=num(r[22]), q=num(r[23])))
    adsets = pd.DataFrame(out)
    h = [k for k, r in enumerate(rows) if r[0] == 'Parent Campaign' and r[2] == 'Ad name'][0]
    out = []
    for r in rows[h + 1:]:
        if r[0] is None or str(r[0]).startswith(('Subtotal', 'Note')): continue
        if str(r[0]).startswith('GRAND'): break
        out.append(dict(camp=r[0], adset=r[1], ad=r[2], status=r[3], spend=num(r[7]), impr=num(r[8]), reach=num(r[9]),
                        qrank=r[11], erank=r[12], crank=r[13], meet=num(r[15]), clicks=num(r[17]), leads=num(r[20]), q=num(r[21])))
    adl = pd.DataFrame(out)
    crm = pd.read_excel(path, "Data_CRM").iloc[:, :8]
    crm.columns = ['date', 'camp', 'adset', 'ad', 'status', 'q', 'm', 'id']
    crm = crm.dropna(subset=['date']); crm['date'] = pd.to_datetime(crm['date']); crm['day'] = crm.date.dt.day
    def grp(r):
        if r.q == 1: return 'Qualified'
        s = r.status
        if s in ('DNP', 'Not Reachable', 'Invalid/Wrong Number'): return 'Unreachable'
        if s == 'Junk Lead': return 'Junk'
        if s in ('Not Interested', 'Lost'): return 'Not interested / Lost'
        return 'Open / not yet worked'
    crm['g'] = crm.apply(grp, axis=1)
    return camp, adsets, adl, crm


# ----------------------------------------------------------------------------- QA
def qa_checks(camp, adsets, adl, crm):
    res = []
    add = lambda n, ok, d: res.append((n, bool(ok), d))
    sc, sa, sd = camp.spend.sum(), adsets.spend.sum(), adl.spend.sum()
    add("Spend reconciles: campaign = ad set = ad (±₹1)", abs(sc - sa) < 1 and abs(sc - sd) < 1, f"{sc:,.2f} / {sa:,.2f} / {sd:,.2f}")
    add("CRM leads: campaign table = Data_CRM rows", int(camp.leads.sum()) == len(crm), f"{int(camp.leads.sum())} vs {len(crm)}")
    add("CRM leads: ad table = Data_CRM rows", int(adl.leads.sum()) == len(crm), f"{int(adl.leads.sum())} vs {len(crm)}")
    add("Qualified: campaign table = CRM flag", int(camp.q.sum()) == int(crm.q.sum()), f"{int(camp.q.sum())} vs {int(crm.q.sum())}")
    add("No duplicate Social Lead IDs", crm.id.duplicated().sum() == 0, f"{int(crm.id.duplicated().sum())} duplicates")
    add("No unmapped leads", (crm.camp.astype(str).str.contains('Not mapped')).sum() == 0, f"{int(crm.camp.astype(str).str.contains('Not mapped').sum())} unmapped")
    add("Meetings: campaign table >= ad table", camp.meet.sum() >= adl.meet.sum(), f"{int(camp.meet.sum())} vs {int(adl.meet.sum())}")
    return res


# ----------------------------------------------------------------------------- excel helpers
def title(ws, t, sub=None):
    ws["A1"] = t; ws["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    if sub: ws["A2"] = sub; ws["A2"].font = Font(name="Arial", size=9, italic=True, color="667085")
    ws.sheet_view.showGridLines = False

def hdr(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c); x.fill = HF; x.font = HFont; x.border = box
        x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w

def put(ws, r, c, v, fmt=None, font=base, fill=None, wrap=False):
    x = ws.cell(r, c, v); x.font = font; x.border = box
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if wrap: x.alignment = Alignment(wrap_text=True, vertical="top")
    return x

def funnel_sheet(ws, df, keycols, keyheads, first_row, total_label):
    nk = len(keycols)
    heads = keyheads + ["Delivery", "Spend (₹)", "Impressions", "Reach", "Link clicks", "CRM leads", "Qualified", "Meetings",
        "CPL (₹ per CRM lead)", "Qualified %", "CPQL (₹)", "Lead→Meeting %", "Qualified→Meeting %", "Cost per Meeting (₹)",
        "CTR", "CPC (₹)", "CPM (₹)", "Frequency", "% of spend", "% of meetings", "Efficiency index (acct cost/mtg ÷ row)"]
    hdr(ws, first_row, heads)
    c0 = nk + 1
    cS, cI, cR, cC, cLd, cQ, cM = [c0 + i for i in range(1, 8)]
    r = first_row + 1; start = r
    for _, x in df.iterrows():
        for j, k in enumerate(keycols): put(ws, r, j + 1, x[k])
        put(ws, r, c0, x.get('status'))
        for off, k in enumerate(['spend', 'impr', 'reach', 'clicks', 'leads', 'q', 'meet']):
            put(ws, r, c0 + 1 + off, x[k], INR if k == 'spend' else INT)
        r += 1
    end = r - 1; tr = r
    put(ws, tr, 1, total_label, font=bold, fill=TOT)
    for j in range(2, nk + 2): put(ws, tr, j, None, fill=TOT)
    for col in (cS, cI, cC, cLd, cQ, cM):
        put(ws, tr, col, f"=SUM({L(col)}{start}:{L(col)}{end})", INR if col == cS else INT, bold, TOT)
    put(ws, tr, cR, None, fill=TOT)
    fcol = cM + 1
    def formulas(rr, isT):
        S, I, R, C, Ld, Q, M = [f"{L(c)}{rr}" for c in (cS, cI, cR, cC, cLd, cQ, cM)]
        f = [(f'=IFERROR({S}/{Ld},"n/a")', INR), (f'=IFERROR({Q}/{Ld},"n/a")', PCT), (f'=IFERROR({S}/{Q},"n/a")', INR),
             (f'=IFERROR({M}/{Ld},"n/a")', PCT), (f'=IFERROR({M}/{Q},"n/a")', PCT),
             (f'=IF(AND({M}>0,{S}>0),{S}/{M},"n/a")', INR), (f'=IFERROR({C}/{I},"n/a")', '0.00%'),
             (f'=IFERROR({S}/{C},"n/a")', INR), (f'=IFERROR({S}/{I}*1000,"n/a")', INR)]
        f.append((f'=IF({R}="","n/a",IFERROR({I}/{R},"n/a"))' if not isT else '="n/a"', FRQ))
        f.append((f'=IFERROR({S}/${L(cS)}${tr},"n/a")', PCT)); f.append((f'=IFERROR({M}/${L(cM)}${tr},"n/a")', PCT))
        f.append((f'=IFERROR(IF(AND({M}>0,{S}>0),(${L(fcol+5)}${tr})/({S}/{M}),"n/a"),"n/a")' if not isT else '="1.00"', FRQ))
        return f
    for rr in range(start, end + 1):
        for k, (fx, fm) in enumerate(formulas(rr, False)): put(ws, rr, fcol + k, fx, fm)
    for k, (fx, fm) in enumerate(formulas(tr, True)): put(ws, tr, fcol + k, fx, fm, bold, TOT)
    ws.cell(tr, fcol + 12).value = 1; ws.cell(tr, fcol + 12).number_format = FRQ
    ws.freeze_panes = ws.cell(first_row + 1, nk + 1)
    return start, end, tr, dict(S=cS, I=cI, R=cR, C=cC, Ld=cLd, Q=cQ, M=cM, F=fcol)


# ----------------------------------------------------------------------------- main build
def build(path, outdir, history_path, targets_path):
    camp, adsets, adl, crm = load(path)
    qa = qa_checks(camp, adsets, adl, crm)
    if camp.leads.sum() == 0 or camp.spend.sum() == 0 or len(crm) == 0:
        print("STOP: the report has no calculated values (open it in Excel, let it recalculate, save, and re-upload) "
              "or it is empty. Failed checks: " + "; ".join(n for n, ok, d in qa if not ok))
        return 3
    end = crm.date.max(); ndays = int(end.day); cut = max(ndays - 3, 1)
    label = end.strftime("%d %b %Y"); stamp = end.strftime("%Y-%m-%d")
    mature = crm[crm.day <= cut]
    targets = dict(DEFAULT_TARGETS)
    if targets_path and os.path.exists(targets_path): targets.update(json.load(open(targets_path)))
    tot_sp = camp.spend.sum()

    wb = Workbook()
    # ---------------- Funnel_Campaign
    wsC = wb.active; wsC.title = "Funnel_Campaign"
    title(wsC, f"Campaign-wise full-funnel scorecard — MTD 1–{label}", "Spend/reach/clicks from Meta; leads, qualified, meetings from CRM via LeadChain. All rates are live formulas.")
    cdf = camp[(camp.spend > 0) | (camp.meet > 0)].copy().sort_values('spend', ascending=False)
    cs, ce, ct, cc = funnel_sheet(wsC, cdf, ['camp'], ["Campaign"], 4, "TOTAL")
    for i, w in enumerate([46, 13, 13, 13, 12, 11, 10, 10, 10, 11, 10, 10, 12, 12, 13, 9, 9, 9, 10, 10, 10, 14], 1): wsC.column_dimensions[L(i)].width = w
    names = list(cdf.camp)
    # best / worst campaign by cost per meeting (spend >= 5% of total, meetings >= 5)
    elig = cdf[(cdf.spend >= 0.05 * tot_sp) & (cdf.meet >= 5)].copy()
    elig['cpm_'] = elig.spend / elig.meet
    best = elig.sort_values('cpm_').iloc[0] if len(elig) else None
    worst = elig.sort_values('cpm_').iloc[-1] if len(elig) > 1 else None
    if best is not None and worst is not None and best.camp != worst.camp:
        r = ct + 3
        wsC.cell(r, 1, "Is the best campaign really better than the weakest? (two-proportion z-test on Qualified %)").font = bold
        iW = cs + names.index(worst.camp); iB = cs + names.index(best.camp)
        Ld, Q, M = L(cc['Ld']), L(cc['Q']), L(cc['M'])
        lab = [(f"Weakest: {worst.camp} — qualified %", f"={Q}{iW}/{Ld}{iW}", PCT), (f"Best: {best.camp} — qualified %", f"={Q}{iB}/{Ld}{iB}", PCT),
               ("Pooled %", f"=({Q}{iW}+{Q}{iB})/({Ld}{iW}+{Ld}{iB})", PCT),
               ("z-score", f"=(B{r+2}-B{r+1})/SQRT(B{r+3}*(1-B{r+3})*(1/{Ld}{iW}+1/{Ld}{iB}))", '0.00'),
               ("Verdict", f'=IF(ABS(B{r+4})>=1.96,"Statistically significant (95%) — gap is real, not noise","Not significant — could be noise")', None),
               ("Weakest qualified→meeting %", f"={M}{iW}/{Q}{iW}", PCT), ("Best qualified→meeting %", f"={M}{iB}/{Q}{iB}", PCT)]
        for k, (t, f, fm) in enumerate(lab, 1): put(wsC, r + k, 1, t); put(wsC, r + k, 2, f, fm)

    # ---------------- Funnel_AdSet
    wsA = wb.create_sheet("Funnel_AdSet")
    title(wsA, f"Ad set-wise full-funnel scorecard — MTD 1–{label}", "Frequency = impressions ÷ reach as reported by Meta for the period.")
    adf = adsets[(adsets.spend > 0) | (adsets.meet > 0)].copy().sort_values('spend', ascending=False)
    a_s, a_e, a_t, ac = funnel_sheet(wsA, adf, ['adset', 'camp'], ["Ad set", "Parent campaign"], 4, "TOTAL")
    for i, w in enumerate([44, 40, 13, 13, 13, 12, 11, 10, 10, 10, 11, 10, 10, 12, 12, 13, 9, 9, 9, 10, 10, 10, 14], 1): wsA.column_dimensions[L(i)].width = w

    # ---------------- Funnel_Ad
    wsD = wb.create_sheet("Funnel_Ad")
    title(wsD, f"Ad-wise scorecard with verdicts — MTD 1–{label}", "Each row is Ad set × Ad. Edit the blue thresholds to change verdicts.")
    th = [("Min spend to judge an ad (₹)", 25000), ("'Scale' if cost/meeting ≤ this × account average", 0.75),
          ("'Reduce/fix' if cost/meeting ≥ this × account average", 1.5), ("Min meetings to call 'Scale'", 5)]
    for k, (t, v) in enumerate(th): wsD.cell(3 + k, 1, t).font = base; wsD.cell(3 + k, 2, v).font = inp
    ddf = adl[adl.spend > 0].copy().sort_values('spend', ascending=False)
    d_s, d_e, d_t, dc = funnel_sheet(wsD, ddf, ['ad', 'adset', 'camp'], ["Ad", "Ad set", "Parent campaign"], 8, "TOTAL (ads with spend)")
    ex = dc['F'] + 13
    for k, h in enumerate(["Qualified % — 95% low", "Qualified % — 95% high", "Sample confidence", "Verdict", "Meta Quality rank", "Meta Engagement rank", "Meta Conversion rank"]):
        x = wsD.cell(8, ex + k, h); x.fill = HF; x.font = HFont; x.border = box; x.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    acct_cpm = ddf.spend.sum() / max(ddf.meet.sum(), 1)
    verdicts = []
    for i, (_, x) in enumerate(ddf.iterrows()):
        rr = d_s + i
        Ld_, Q_, M_, S_ = (f"{L(dc[k])}{rr}" for k in ('Ld', 'Q', 'M', 'S'))
        p = f"({Q_}/{Ld_})"; z = "1.96"
        put(wsD, rr, ex, f'=IF({Ld_}=0,"n/a",({p}+{z}^2/(2*{Ld_})-{z}*SQRT({p}*(1-{p})/{Ld_}+{z}^2/(4*{Ld_}^2)))/(1+{z}^2/{Ld_}))', PCT)
        put(wsD, rr, ex + 1, f'=IF({Ld_}=0,"n/a",({p}+{z}^2/(2*{Ld_})+{z}*SQRT({p}*(1-{p})/{Ld_}+{z}^2/(4*{Ld_}^2)))/(1+{z}^2/{Ld_}))', PCT)
        put(wsD, rr, ex + 2, f'=IF({Ld_}>=100,"High",IF({Ld_}>=40,"Medium","Low"))')
        cm = f"{L(dc['F']+5)}{rr}"; avg = f"${L(dc['F']+5)}${d_t}"
        put(wsD, rr, ex + 3, f'=IF({Ld_}=0,"n/a (not a lead-gen ad)",IF({S_}<$B$3,"Too early (low spend)",IF({M_}=0,"Pause / fix — no meetings",IF(AND({cm}<=$B$4*{avg},{M_}>=$B$6),"SCALE",IF({cm}>=$B$5*{avg},"REDUCE / fix","Hold & watch")))))', font=bold)
        cl = lambda s: (s or '-').replace(' of ads', '')
        put(wsD, rr, ex + 4, cl(x.qrank)); put(wsD, rr, ex + 5, cl(x.erank)); put(wsD, rr, ex + 6, cl(x.crank))
        # python mirror of the verdict for narrative
        v = "n/a"
        if x.leads > 0:
            if x.spend < 25000: v = "Too early"
            elif x.meet == 0: v = "Pause / fix"
            else:
                c_ = x.spend / x.meet
                v = "SCALE" if (c_ <= 0.75 * acct_cpm and x.meet >= 5) else ("REDUCE / fix" if c_ >= 1.5 * acct_cpm else "Hold & watch")
        verdicts.append((x.ad, x.adset, x.camp, x.spend, int(x.leads), int(x.q), int(x.meet), v))
    for i, w in enumerate([34, 40, 40, 13, 13, 13, 12, 11, 10, 10, 10, 11, 10, 10, 12, 12, 13, 9, 9, 9, 10, 10, 10, 14, 12, 12, 12, 26, 22, 22, 22], 1): wsD.column_dimensions[L(i)].width = w
    for key, color in (("SCALE", "C6EFCE"), ("REDUCE", "FFC7CE"), ("Pause", "FFC7CE")):
        wsD.conditional_formatting.add(f"{L(ex+3)}{d_s}:{L(ex+3)}{d_e}", FormulaRule(formula=[f'LEFT({L(ex+3)}{d_s},{len(key)})="{key}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
    wsD.freeze_panes = wsD.cell(9, 4)
    vdf = pd.DataFrame(verdicts, columns=['ad', 'adset', 'camp', 'spend', 'leads', 'q', 'meet', 'verdict'])
    vdf['cpm'] = np.where(vdf.meet > 0, vdf.spend / vdf.meet.replace(0, np.nan), np.nan)

    # ---------------- Lead_Quality
    wsL = wb.create_sheet("Lead_Quality")
    title(wsL, "Lead quality & contactability — why leads do or don't convert", "Qualified = CRM Qualified flag. Unreachable = DNP / Not Reachable / Invalid number.")
    def qtable(top, keys, heads):
        g = crm.groupby(keys + ['g']).size().unstack(fill_value=0).reindex(columns=GROUPS, fill_value=0).reset_index()
        g['tot'] = g[GROUPS].sum(axis=1); g = g.sort_values('tot', ascending=False)
        hdr(wsL, top, heads + ["Leads"] + GROUPS + ["Qualified %", "Open %", "Unreachable %", "Junk %", "Not interested %", "Reachable-and-engaged %"])
        nk = len(keys); r = top + 1; s = r
        for _, x in g.iterrows():
            for j, k in enumerate(keys): put(wsL, r, j + 1, x[k])
            put(wsL, r, nk + 1, f"=SUM({L(nk+2)}{r}:{L(nk+6)}{r})", INT)
            for j, gname in enumerate(GROUPS): put(wsL, r, nk + 2 + j, int(x[gname]), INT)
            for j in range(5): put(wsL, r, nk + 7 + j, f"=IFERROR({L(nk+2+j)}{r}/{L(nk+1)}{r},0)", PCT)
            put(wsL, r, nk + 12, f"=IFERROR(({L(nk+2)}{r}+{L(nk+3)}{r})/{L(nk+1)}{r},0)", PCT); r += 1
        put(wsL, r, 1, "TOTAL", font=bold, fill=TOT)
        for j in range(2, nk + 1): put(wsL, r, j, None, fill=TOT)
        for c in range(nk + 1, nk + 7): put(wsL, r, c, f"=SUM({L(c)}{s}:{L(c)}{r-1})", INT, bold, TOT)
        for j in range(5): put(wsL, r, nk + 7 + j, f"=IFERROR({L(nk+2+j)}{r}/{L(nk+1)}{r},0)", PCT, bold, TOT)
        put(wsL, r, nk + 12, f"=IFERROR(({L(nk+2)}{r}+{L(nk+3)}{r})/{L(nk+1)}{r},0)", PCT, bold, TOT)
        return r
    r1 = qtable(4, ['camp'], ["Campaign"])
    wsL.cell(r1 + 2, 1, "By ad set × ad").font = bold
    r2 = qtable(r1 + 3, ['adset', 'ad', 'camp'], ["Ad set", "Ad", "Campaign"])
    wsL.cell(r2 + 2, 1, "Raw CRM status counts").font = bold
    rs = r2 + 3; hdr(wsL, rs, ["Lead status", "Leads", "% of leads"]); sc = crm.status.value_counts()
    for k, (s_, n) in enumerate(sc.items(), 1):
        put(wsL, rs + k, 1, s_); put(wsL, rs + k, 2, int(n), INT); put(wsL, rs + k, 3, f"=B{rs+k}/SUM($B${rs+1}:$B${rs+len(sc)})", PCT)
    for i, w in enumerate([46, 44, 44, 11, 11, 13, 13, 11, 11, 11, 11, 11, 11, 13, 13], 1): wsL.column_dimensions[L(i)].width = w

    # ---------------- Trends
    wsT = wb.create_sheet("Trends")
    title(wsT, "Daily trend by lead-creation date — volume, quality and maturity", "'Open %' = leads sales has not worked yet; recent days look worse than they will end up.")
    dates = crm.date.dt.normalize()
    daily = crm.groupby(dates).agg(n=('id', 'count'), q=('q', 'sum')).reset_index().rename(columns={'date': 'd'})
    op = crm[crm.g == 'Open / not yet worked'].groupby(dates).size(); un = crm[crm.g == 'Unreachable'].groupby(dates).size()
    hdr(wsT, 4, ["Lead date", "Leads", "Qualified", "Qualified %", "Open (not yet worked)", "Open %", "Unreachable", "Unreachable %",
                 "7-day rolling leads (avg/day)", "7-day rolling qualified %", "Status of data"], [46, 10, 10, 11, 13, 10, 12, 12, 15, 15, 22])
    for i, x in daily.iterrows():
        rr = 5 + i; d = x['d']
        put(wsT, rr, 1, d.to_pydatetime(), 'dd-mmm'); put(wsT, rr, 2, int(x.n), INT); put(wsT, rr, 3, int(x.q), INT)
        put(wsT, rr, 4, f"=C{rr}/B{rr}", PCT); put(wsT, rr, 5, int(op.get(d, 0)), INT); put(wsT, rr, 6, f"=E{rr}/B{rr}", PCT)
        put(wsT, rr, 7, int(un.get(d, 0)), INT); put(wsT, rr, 8, f"=G{rr}/B{rr}", PCT)
        lo = max(5, rr - 6)
        put(wsT, rr, 9, f"=AVERAGE(B{lo}:B{rr})", '0.0'); put(wsT, rr, 10, f"=SUM(C{lo}:C{rr})/SUM(B{lo}:B{rr})", PCT)
        put(wsT, rr, 11, f'=IF(F{rr}>0.3,"Immature — too fresh",IF(F{rr}>0.15,"Maturing","Mature"))')
    te = 5 + len(daily) - 1
    ch = BarChart(); ch.title = "Daily CRM leads and qualified %"; ch.y_axis.title = "Leads"; ch.height = 8.5; ch.width = 24
    ch.add_data(Reference(wsT, min_col=2, min_row=4, max_row=te), titles_from_data=True); ch.set_categories(Reference(wsT, min_col=1, min_row=5, max_row=te))
    ln = LineChart(); ln.add_data(Reference(wsT, min_col=4, min_row=4, max_row=te), titles_from_data=True)
    ln.y_axis.axId = 200; ln.y_axis.title = "Qualified %"; ln.y_axis.crosses = "max"; ln.y_axis.number_format = '0%'
    ch += ln; wsT.add_chart(ch, "M4")
    pr = te + 3
    edges = np.linspace(1, cut + 1, 5).astype(int)
    periods = [(f"{edges[i]}–{edges[i+1]-1} {end.strftime('%b')}", edges[i], edges[i + 1] - 1) for i in range(4) if edges[i + 1] - 1 >= edges[i]]
    wsT.cell(pr, 1, f"Qualified rate by lead period — only leads up to day {cut} (≥3 days old) so immature leads don't distort").font = bold
    hdr(wsT, pr + 1, ["Period (lead date)", "Leads", "Qualified", "Qualified %", "Unreachable %", "Junk %"])
    period_stats = []
    for k, (nm, a, b) in enumerate(periods):
        s = crm[(crm.day >= a) & (crm.day <= b)]; rr = pr + 2 + k
        put(wsT, rr, 1, nm); put(wsT, rr, 2, len(s), INT); put(wsT, rr, 3, int(s.q.sum()), INT); put(wsT, rr, 4, f"=C{rr}/B{rr}", PCT)
        put(wsT, rr, 5, float((s.g == 'Unreachable').mean()), PCT); put(wsT, rr, 6, float((s.g == 'Junk').mean()), PCT)
        period_stats.append((nm, len(s), float(s.q.mean()), float((s.g == 'Unreachable').mean())))
    mr = pr + 2 + len(periods) + 2
    mid = cut // 2
    wsT.cell(mr, 1, f"Ad set mix shift — first half (days 1–{mid}) vs second half (days {mid+1}–{cut}), mature leads only").font = bold
    hdr(wsT, mr + 1, ["Ad set", f"Leads 1–{mid}", f"Qualified % 1–{mid}", f"Unreachable % 1–{mid}", f"Leads {mid+1}–{cut}", f"Qualified % {mid+1}–{cut}", f"Unreachable % {mid+1}–{cut}"])
    pre = mature[mature.day <= mid]; post = mature[mature.day > mid]
    sets = list(mature.adset.value_counts().index)
    for k, sname in enumerate(sets):
        rr = mr + 2 + k; a_ = pre[pre.adset == sname]; b_ = post[post.adset == sname]
        put(wsT, rr, 1, sname); put(wsT, rr, 2, len(a_), INT); put(wsT, rr, 3, float(a_.q.mean()) if len(a_) else None, PCT); put(wsT, rr, 4, float((a_.g == 'Unreachable').mean()) if len(a_) else None, PCT)
        put(wsT, rr, 5, len(b_), INT); put(wsT, rr, 6, float(b_.q.mean()) if len(b_) else None, PCT); put(wsT, rr, 7, float((b_.g == 'Unreachable').mean()) if len(b_) else None, PCT)
    rr = mr + 2 + len(sets)
    put(wsT, rr, 1, "ALL", font=bold, fill=TOT)
    for c, v, f in ((2, len(pre), INT), (3, float(pre.q.mean()), PCT), (4, float((pre.g == 'Unreachable').mean()), PCT), (5, len(post), INT), (6, float(post.q.mean()), PCT), (7, float((post.g == 'Unreachable').mean()), PCT)):
        put(wsT, rr, c, v, f, bold, TOT)

    # ---------------- Budget_Scenarios
    wsB = wb.create_sheet("Budget_Scenarios")
    title(wsB, "Budget pacing & reallocation scenario", "Blue cells are inputs. Scenario uses MTD average efficiency — real marginal returns will be lower, hence the haircut.")
    hdr(wsB, 4, ["Campaign", "Daily budget in export (₹)", "MTD spend (₹)", f"Avg spend/day (₹, ÷{ndays})", "Budget utilisation", "Read-out"], [58, 22, 22, 22, 22, 62])
    bud = camp[camp.spend > 0].copy(); util_notes = []
    for k, x in enumerate(bud.itertuples()):
        rr = 5 + k; isn = isinstance(x.budget, (int, float))
        put(wsB, rr, 1, x.camp); put(wsB, rr, 2, x.budget if isn else "ad set budget", INR if isn else None); put(wsB, rr, 3, x.spend, INR)
        put(wsB, rr, 4, f"=C{rr}/{ndays}", INR); put(wsB, rr, 5, f'=IFERROR(D{rr}/B{rr},"n/a")', PCT)
        put(wsB, rr, 6, f'=IF(ISNUMBER(E{rr}),IF(E{rr}<0.75,"Under-delivering vs current budget — budget raised recently, or limited by audience/bid",IF(E{rr}>1.1,"Over budget average (budget lowered, or CBO shifted)","Spending ≈ budget")),"n/a")')
        if isn and x.budget > 0: util_notes.append((x.camp, x.spend / ndays / x.budget, x.budget))
    sc0 = 5 + len(bud) + 2
    if best is not None and worst is not None and best.camp != worst.camp:
        wsB.cell(sc0, 1, f"Scenario: move budget from '{worst.camp}' → '{best.camp}' (MTD average efficiency, diminishing-returns haircut)").font = bold
        put(wsB, sc0 + 1, 1, "Amount moved (₹, MTD-equivalent)"); put(wsB, sc0 + 1, 2, round(0.15 * worst.spend, -3), INR, inp)
        iW = cs + names.index(worst.camp); iB = cs + names.index(best.camp)
        fm = lambda col, row: f"Funnel_Campaign!{L(cc[col])}{row}"
        for k, (t, f) in enumerate([("Weakest: qualified per ₹", f"={fm('Q',iW)}/{fm('S',iW)}"), ("Weakest: meetings per ₹", f"={fm('M',iW)}/{fm('S',iW)}"),
                                     ("Best: qualified per ₹", f"={fm('Q',iB)}/{fm('S',iB)}"), ("Best: meetings per ₹", f"={fm('M',iB)}/{fm('S',iB)}")], 2):
            put(wsB, sc0 + k, 1, t); put(wsB, sc0 + k, 2, f, '0.000000')
        hdr(wsB, sc0 + 7, ["Haircut on best campaign's efficiency", "Qualified lost (weakest)", "Meetings lost (weakest)", "Qualified gained (best)", "Meetings gained (best)", "Net change"])
        for k, h in enumerate([0, 0.25, 0.5, 0.75]):
            rr = sc0 + 8 + k; put(wsB, rr, 1, h, PCT, inp)
            put(wsB, rr, 2, f"=$B${sc0+1}*$B${sc0+2}", '0.0'); put(wsB, rr, 3, f"=$B${sc0+1}*$B${sc0+3}", '0.0')
            put(wsB, rr, 4, f"=$B${sc0+1}*$B${sc0+4}*(1-A{rr})", '0.0'); put(wsB, rr, 5, f"=$B${sc0+1}*$B${sc0+5}*(1-A{rr})", '0.0')
            put(wsB, rr, 6, f'=TEXT(D{rr}-B{rr},"+0.0;-0.0")&" qualified  |  "&TEXT(E{rr}-C{rr},"+0.0;-0.0")&" meetings"')
        put(wsB, sc0 + 13, 1, "Break-even haircut — meetings"); put(wsB, sc0 + 13, 2, f"=1-B{sc0+3}/B{sc0+5}", PCT)
        put(wsB, sc0 + 14, 1, "Break-even haircut — qualified leads"); put(wsB, sc0 + 14, 2, f"=1-B{sc0+2}/B{sc0+4}", PCT)
        wsB.cell(sc0 + 16, 1, "Directional only. Test a 10–15% shift for 7 days and measure real marginal cost before committing.").font = Font(name="Arial", size=9, italic=True)

    # ---------------- Daily_Delta + history
    delta_rows = None; hist = None
    snap = []
    for lvl, df_, key in (("campaign", camp, lambda x: x.camp), ("adset", adsets, lambda x: f"{x.camp} || {x.adset}"), ("ad", adl, lambda x: f"{x.camp} || {x.adset} || {x.ad}")):
        for x in df_.itertuples():
            if x.spend > 0 or x.leads > 0 or x.meet > 0:
                snap.append(dict(date=stamp, level=lvl, key=key(x), spend=round(x.spend, 2), impr=x.impr, reach=x.reach, clicks=x.clicks, leads=x.leads, qualified=x.q, meetings=x.meet))
    snap = pd.DataFrame(snap)
    if history_path:
        os.makedirs(os.path.dirname(history_path) or ".", exist_ok=True)
        hist = pd.read_csv(history_path) if os.path.exists(history_path) else pd.DataFrame(columns=snap.columns)
        hist = hist[hist.date != stamp]
        prev_dates = sorted(d for d in hist.date.unique() if d < stamp and d[:7] == stamp[:7])
        if prev_dates:
            pd_ = prev_dates[-1]; pv = hist[(hist.date == pd_) & (hist.level == 'campaign')].set_index('key')
            cu = snap[snap.level == 'campaign'].set_index('key')
            delta_rows = []
            for k_ in cu.index:
                o = pv.loc[k_] if k_ in pv.index else None
                g = lambda c: float(cu.loc[k_, c] - (o[c] if o is not None else 0))
                delta_rows.append([k_, g('spend'), g('impr'), g('clicks'), g('leads'), g('qualified'), g('meetings')])
            wsX = wb.create_sheet("Daily_Delta")
            title(wsX, f"True daily change: {stamp} MTD minus {pd_} MTD", "Meta exports are cumulative, so the difference between consecutive days gives the day's spend, impressions and clicks. Leads/qualified/meetings can also move because CRM statuses change.")
            hdr(wsX, 4, ["Campaign", "Day spend (₹)", "Day impressions", "Day link clicks", "Day CRM leads (net)", "Day qualified (net)", "Day meetings (net)", "Day CPL (₹)", "Day CTR"], [58, 16, 16, 16, 16, 16, 16, 14, 12])
            for k, row in enumerate(delta_rows):
                rr = 5 + k
                put(wsX, rr, 1, row[0]); put(wsX, rr, 2, row[1], INR)
                for j, f in enumerate((INT,) * 5): put(wsX, rr, 3 + j, row[2 + j], f)
                put(wsX, rr, 8, f'=IFERROR(B{rr}/E{rr},"n/a")', INR); put(wsX, rr, 9, f'=IFERROR(D{rr}/C{rr},"n/a")', '0.00%')
        out = pd.concat([hist, snap], ignore_index=True)
        out.to_csv(history_path, index=False)

    # ---------------- Executive summary (auto-generated)
    cA = lambda c, r: f"Funnel_Campaign!{L(cc[c])}{r}"
    A_leads, A_q, A_m = int(camp.leads.sum()), int(camp.q.sum()), int(camp.meet.sum())
    cost_mtg = tot_sp / A_m if A_m else float('nan')
    wsE = wb.create_sheet("Executive_Summary", 0)
    title(wsE, f"DAEWOO × SHOOTORDER — Deep Performance Analysis | MTD 1–{label}", "Auto-generated from your daily Campaign Performance Report. Figures in INR. Every rate in the detail tabs is a live formula.")
    kp = [("MTD spend", f"=Funnel_Campaign!{L(cc['S'])}{ct}", INR), ("CRM leads", f"=Funnel_Campaign!{L(cc['Ld'])}{ct}", INT), ("CPL", f"=Funnel_Campaign!{L(cc['F'])}{ct}", INR),
          ("Qualified leads", f"=Funnel_Campaign!{L(cc['Q'])}{ct}", INT), ("Qualified %", f"=Funnel_Campaign!{L(cc['F']+1)}{ct}", PCT), ("CPQL", f"=Funnel_Campaign!{L(cc['F']+2)}{ct}", INR),
          ("Meetings (mapped)", f"=Funnel_Campaign!{L(cc['M'])}{ct}", INT), ("Lead→Meeting %", f"=Funnel_Campaign!{L(cc['F']+3)}{ct}", PCT), ("Cost per Meeting", f"=Funnel_Campaign!{L(cc['F']+5)}{ct}", INR)]
    for k, (t, f, fm_) in enumerate(kp):
        c = 1 + k; x = wsE.cell(4, c, t); x.fill = HF; x.font = HFont; x.alignment = Alignment(horizontal="center", wrap_text=True); x.border = box
        y = wsE.cell(5, c, f); y.number_format = fm_; y.font = Font(name="Arial", size=13, bold=True, color=NAVY); y.alignment = Alignment(horizontal="center"); y.border = box; y.fill = PatternFill("solid", fgColor=GREY)
    wsE.row_dimensions[5].height = 26
    for c in range(1, 10): wsE.column_dimensions[L(c)].width = 19

    F = []  # findings (headline, detail)
    ctab = cdf.assign(cpl=lambda d: d.spend / d.leads.replace(0, np.nan), qp=lambda d: d.q / d.leads.replace(0, np.nan),
                      q2m=lambda d: d.meet / d.q.replace(0, np.nan), cpmtg=lambda d: d.spend / d.meet.replace(0, np.nan))
    if best is not None and worst is not None and best.camp != worst.camp:
        b_ = ctab[ctab.camp == best.camp].iloc[0]; w_ = ctab[ctab.camp == worst.camp].iloc[0]
        F.append(("CPL is the wrong yardstick — cost per meeting is the real one.",
                  f"'{best.camp}' costs {inr(b_.cpl)} per lead vs {inr(w_.cpl)} for '{worst.camp}', but qualifies {pc(b_.qp)} vs {pc(w_.qp)} and converts {pc(b_.q2m)} of qualified leads to meetings vs {pc(w_.q2m)}. Cost per meeting: {inr(b_.cpmtg)} vs {inr(w_.cpmtg)}."))
        F.append(("Spend and results are mismatched.",
                  f"'{worst.camp}' takes {pc(w_.spend/tot_sp)} of spend and delivers {pc(w_.meet/A_m)} of meetings. '{best.camp}' takes {pc(b_.spend/tot_sp)} of spend and delivers {pc(b_.meet/A_m)} of meetings."))
    if len(period_stats) >= 2:
        s0, s1 = period_stats[0], period_stats[-1]
        trend = " → ".join(f"{n}: {pc(q)}" for n, _, q, _ in period_stats)
        verb = "fell" if s1[2] < s0[2] - 0.02 else ("rose" if s1[2] > s0[2] + 0.02 else "held steady")
        F.append((f"Lead quality {verb} across the month (mature leads only).",
                  f"Qualified %: {trend}. Unreachable share moved from {pc(s0[3])} to {pc(s1[3])}. Daily lead volume averaged {daily.n[:7].mean():.0f}/day in the first week vs {daily.n[-10:-3].mean():.0f}/day in the last mature week."))
    vv = vdf[vdf.leads > 0]
    sc_ = vv[vv.verdict == "SCALE"].sort_values('cpm'); rd_ = vv[vv.verdict.isin(["REDUCE / fix", "Pause / fix"])].sort_values('spend', ascending=False)
    if len(sc_):
        x = sc_.iloc[0]; F.append(("Best-performing ad(s) to scale.", f"'{x.ad}' in '{x.adset}': {int(x.meet)} meetings at {inr(x.cpm)} each on {inr(x.spend)} spend. Account average across ads is {inr(acct_cpm)} per meeting."))
    if len(rd_):
        x = rd_.iloc[0]; F.append(("Largest drag among ads.", f"'{x.ad}' in '{x.adset}': {inr(x.spend)} spent for {int(x.meet)} meeting(s)" + (f" ({inr(x.cpm)} each)" if x.meet else "") + f"; {len(rd_)} ad(s) are flagged REDUCE / fix on Funnel_Ad."))
    unr = (crm.g == 'Unreachable').mean(); jk = (crm.g == 'Junk').mean(); opn = (crm.g == 'Open / not yet worked').mean()
    F.append(("Contactability is a leak.", f"{pc(unr)} of leads are DNP / not reachable / invalid and {pc(jk)} are junk, so {pc(unr+jk)} never become a real conversation; {pc(opn)} are still open."))
    fresh = crm[crm.day >= ndays - 1]
    F.append((f"Don't judge the last 2 days ({fresh.q.mean()*100:.1f}% qualified) yet.", f"{pc((fresh.g=='Open / not yet worked').mean())} of leads from the last 2 days are still unworked in the CRM. Judge a cohort after ~3 days."))
    for nm, u, b in util_notes:
        if u < 0.75 or u > 1.1:
            F.append((f"Budget pacing: {nm}.", f"Averaging {inr(b*u)}/day vs {inr(b)} daily budget ({pc(u)}). " + ("Under-delivering — raised recently or limited by audience/bid; confirm with a daily export." if u < 0.75 else "Over budget average — budget lowered or CBO shifted.")))
    zs = camp[(camp.spend == 0) & (camp.meet > 0)]
    if len(zs): F.append(("Zero-spend campaigns still delivered meetings.", f"{int(zs.meet.sum())} of {A_m} meetings ({pc(zs.meet.sum()/A_m)}) came from campaigns with no spend this month — older leads converting, so cost per meeting is lagged and blended."))
    wsE.cell(7, 1, "THE FUNNEL IN ONE LINE").font = Font(name="Arial", size=11, bold=True, color=NAVY)
    wsE.cell(8, 1, f"{inr(tot_sp)} spend → {A_leads:,} CRM leads (CPL {inr(tot_sp/A_leads)}) → {A_q} qualified ({pc(A_q/A_leads)}, CPQL {inr(tot_sp/A_q)}) → {A_m} meetings ({pc(A_m/A_leads)} of leads) → {inr(cost_mtg)} per meeting. Won/revenue is not in the data, so ROAS cannot be measured.").font = base
    wsE.merge_cells("A8:I8"); wsE["A8"].alignment = Alignment(wrap_text=True, vertical="top"); wsE.row_dimensions[8].height = 32
    wsE.cell(10, 1, "KEY FINDINGS").font = Font(name="Arial", size=11, bold=True, color=NAVY)
    r = 11
    for n, (h, t) in enumerate(F, 1):
        wsE.cell(r, 1, n).font = Font(name="Arial", size=11, bold=True, color="FFFFFF"); wsE.cell(r, 1).fill = HF; wsE.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        wsE.cell(r, 2, h).font = bold; wsE.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        wsE.cell(r + 1, 2, t).font = base; wsE.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=9)
        wsE.cell(r + 1, 2).alignment = Alignment(wrap_text=True, vertical="top"); wsE.row_dimensions[r + 1].height = max(32, 15 * (len(t) // 125 + 1)); r += 2

    # ---------------- Action_Plan (rule-driven + standing actions)
    wsP = wb.create_sheet("Action_Plan")
    title(wsP, "Action plan — what to do, why, and how we will know it worked", "Rule-driven from today's data (ad verdicts, budget pacing) plus standing actions. Edit thresholds on Funnel_Ad.")
    hdr(wsP, 4, ["#", "Priority", "Action", "Evidence", "Success measure", "Review after"], [5, 10, 60, 62, 42, 14])
    acts = []
    for x in sc_.head(3).itertuples():
        acts.append(("P1", f"Scale '{x.ad}' in '{x.adset}' by ~20% every 3 days; clone the creative into other ad sets.", f"{int(x.meet)} meetings at {inr(x.cpm)} vs account {inr(acct_cpm)}; spend {inr(x.spend)}.", "Cost per meeting stays ≤ 1.2× current while spend grows", "7 days"))
    for x in rd_.head(3).itertuples():
        acts.append(("P1", f"Pause or replace '{x.ad}' in '{x.adset}'; reallocate its budget to the scale list.", f"{inr(x.spend)} spent, {int(x.meet)} meeting(s)" + (f" ({inr(x.cpm)} each)" if x.meet else "") + f"; account average {inr(acct_cpm)}.", "Replacement reaches account-average cost per meeting within 150 leads", "10 days"))
    acts += [("P1", "Speed-to-lead SLA: first call within 15 min in working hours, 3 attempts across day-parts, WhatsApp follow-up on DNP.", f"{pc(unr)} of leads unreachable; {pc(opn)} still open.", "Contact rate ≥ 70%; DNP after 3 attempts ≤ 25%", "14 days"),
             ("P2", "Add 1–2 qualifying questions to the instant form on the campaigns with the highest junk / unreachable share.", "See Lead_Quality tab by campaign and ad set.", "Junk ≤ 5%; qualified % up on mature leads", "14 days"),
             ("P2", "Add Won status, deal value, meeting date and BDE to the CRM export.", "Report ends at 'meeting scheduled'; revenue and lag cannot be measured.", "Cost per Won and ROAS reported weekly", "30 days"),
             ("P2", "Export Meta with a 'Day' breakdown (and age / region / placement weekly).", "Daily_Delta (from history) works only after 2+ days of snapshots; breakdowns are missing.", "Daily spend/CPL trend and region/age cuts available", "Next report"),
             ("P3", "Agree KPI targets and fill the blue cells on Measurement_Framework.", "No client-agreed targets in the data.", "All KPIs show On / Off target", "7 days")]
    for k, a in enumerate(acts):
        for j, v in enumerate([k + 1, a[0], a[1], a[2], a[3], a[4]], 1): put(wsP, 5 + k, j, v, font=bold if j == 2 else base, wrap=True)
        wsP.row_dimensions[5 + k].height = 48
    wsP.freeze_panes = "A5"
    recs = [a[1] for a in acts[:4]]
    rr = r + 1
    wsE.cell(rr, 1, "TOP ACTIONS (full list in Action_Plan)").font = Font(name="Arial", size=11, bold=True, color=NAVY)
    for k, a in enumerate(acts[:6], 1):
        wsE.cell(rr + k, 1, k).font = bold; wsE.cell(rr + k, 1).alignment = Alignment(horizontal="center", vertical="top")
        wsE.cell(rr + k, 2, a[1]).font = base; wsE.merge_cells(start_row=rr + k, start_column=2, end_row=rr + k, end_column=9)
        wsE.cell(rr + k, 2).alignment = Alignment(wrap_text=True, vertical="top"); wsE.row_dimensions[rr + k].height = 30

    # ---------------- Measurement_Framework
    wsM = wb.create_sheet("Measurement_Framework")
    title(wsM, "How to measure deeply — KPI tree, definitions and targets", "Level 1 = business outcome; Level 2 = funnel health; Level 3 = diagnostics. Targets come from config/targets.json (blue cells) — overwrite with client-agreed numbers.")
    hdr(wsM, 4, ["Level", "KPI", "Definition / formula", "Source today", "Current (MTD)", "Target", "Check frequency", "Status"], [8, 34, 52, 30, 16, 16, 15, 18])
    T = targets
    ms = [("1", "Cost per Won (CPW)", "Spend ÷ Won deals", "NOT AVAILABLE — need Won status", None, None, "Weekly", INR, True),
          ("1", "ROAS", "Closed deal value ÷ Spend", "NOT AVAILABLE — need deal value", None, None, "Weekly", FRQ, False),
          ("1", "Cost per Meeting", "Spend ÷ Meetings", "Meta + CRM", f"=Funnel_Campaign!{L(cc['F']+5)}{ct}", T['cost_per_meeting'], "Daily", INR, True),
          ("2", "CPQL", "Spend ÷ Qualified leads", "Meta + CRM", f"=Funnel_Campaign!{L(cc['F']+2)}{ct}", T['cpqk'], "Daily", INR, True),
          ("2", "Qualified % (mature leads)", "Qualified ÷ leads older than 3 days", "CRM", float(mature.q.mean()), T['qualified_pct'], "Daily", PCT, False),
          ("2", "Qualified → Meeting %", "Meetings ÷ Qualified", "CRM", f"=Funnel_Campaign!{L(cc['F']+4)}{ct}", T['qual_to_meeting'], "Weekly", PCT, False),
          ("2", "CPL", "Spend ÷ CRM leads", "Meta + CRM", f"=Funnel_Campaign!{L(cc['F'])}{ct}", T['cpl'], "Daily", INR, True),
          ("3", "Contact rate", "(Leads − DNP − Not reachable − Invalid) ÷ Leads", "CRM", float(1 - unr), T['contact_rate'], "Daily", PCT, False),
          ("3", "Junk rate", "Junk ÷ Leads", "CRM", float(jk), T['junk_rate'], "Daily", PCT, True),
          ("3", "Speed to lead (min)", "First call time − lead time", "NOT AVAILABLE — need call log", None, T['speed_to_lead_min'], "Daily", '0', True),
          ("3", "CTR (link)", "Link clicks ÷ Impressions", "Meta", f"=Funnel_Campaign!{L(cc['F']+6)}{ct}", T['ctr'], "Daily", '0.00%', False),
          ("3", "Click → Lead %", "CRM leads ÷ Link clicks", "Meta + CRM", f"=Funnel_Campaign!{L(cc['Ld'])}{ct}/Funnel_Campaign!{L(cc['C'])}{ct}", T['click_to_lead'], "Daily", PCT, False),
          ("3", "Meeting lag (days)", "Meeting date − Lead date", "NOT AVAILABLE — meeting date missing", None, None, "Weekly", '0', True)]
    for k, m_ in enumerate(ms):
        rr = 5 + k
        put(wsM, rr, 1, "L" + m_[0], font=bold); put(wsM, rr, 2, m_[1], font=bold); put(wsM, rr, 3, m_[2], wrap=True); put(wsM, rr, 4, m_[3])
        put(wsM, rr, 5, m_[4], m_[7]); put(wsM, rr, 6, m_[5], m_[7], inp); put(wsM, rr, 7, m_[6])
        op_ = "<=" if m_[8] else ">="
        put(wsM, rr, 8, f'=IF(OR(E{rr}="",F{rr}=""),"Need data/target",IF(E{rr}{op_}F{rr},"On target","Off target"))')
    end_m = 4 + len(ms)
    for txt, col in (("Off target", "FFC7CE"), ("On target", "C6EFCE")):
        wsM.conditional_formatting.add(f"H5:H{end_m}", FormulaRule(formula=[f'H5="{txt}"'], fill=PatternFill("solid", bgColor=col, fgColor=col)))

    # ---------------- QA + Methodology
    wsQ = wb.create_sheet("QA_Checks")
    title(wsQ, "Automatic QA checks run on the input report", "Any FAIL means do not send this report until the source file is fixed.")
    hdr(wsQ, 4, ["Check", "Result", "Detail"], [60, 12, 50])
    for k, (n, ok, d) in enumerate(qa):
        put(wsQ, 5 + k, 1, n); put(wsQ, 5 + k, 2, "PASS" if ok else "FAIL", font=bold); put(wsQ, 5 + k, 3, d)
    wsQ.conditional_formatting.add(f"B5:B{4+len(qa)}", FormulaRule(formula=['B5="FAIL"'], fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")))
    wsN = wb.create_sheet("Methodology_Caveats")
    title(wsN, "Methodology, definitions and caveats")
    notes = [("Source", "All numbers come from your daily 'Campaign Report' workbook (Meta exports + CRM mapped via LeadChain). Nothing is re-sourced."),
             ("Qualified", "CRM Qualified flag (statuses Qualified, Interested, Details Shared, Negotiation, Token Received, Won, Follow-Up, Meeting Scheduled/Conducted/Follow-Up)."),
             ("Lead-quality groups", "Qualified • Open/not yet worked • Unreachable (DNP, Not Reachable, Invalid) • Junk • Not interested/Lost."),
             ("Meetings", "Counted by CRM date; can come from older leads (MeetingDetails), so cost per meeting is a blended, lagged number. Ad-level total can be lower than campaign-level (zero-spend campaigns, unmapped ads)."),
             ("Maturity", "CRM statuses take ~3 days to settle. 'Mature' = lead ≥3 days old. Use cohorts, not raw latest-day rates."),
             ("Daily spend", "Meta exports are cumulative. Daily spend/CTR/CPL appear on Daily_Delta once two consecutive days of snapshots exist in history/daily_snapshot.csv."),
             ("Small samples", "Ads with <40 leads are 'Low' confidence; <₹25,000 spend are 'Too early'. 95% intervals use the Wilson method."),
             ("Meta rankings", "Quality / Engagement / Conversion rankings are Meta's relative scores and often do not predict CRM outcomes."),
             ("Scenario model", "Reallocation uses MTD average efficiency per ₹ with a user-set haircut. Directional, not a forecast."),
             ("Not in data", "Won status, deal value, revenue, call attempts, speed to lead, BDE, region, age, placement.")]
    for k, (a, b_) in enumerate(notes):
        put(wsN, 4 + k, 1, a, font=bold, wrap=True); put(wsN, 4 + k, 2, b_, wrap=True); wsN.row_dimensions[4 + k].height = 40
    wsN.column_dimensions['A'].width = 22; wsN.column_dimensions['B'].width = 130

    order = ["Executive_Summary", "Funnel_Campaign", "Funnel_AdSet", "Funnel_Ad", "Lead_Quality", "Trends", "Budget_Scenarios"] + (["Daily_Delta"] if delta_rows else []) + ["Action_Plan", "Measurement_Framework", "QA_Checks", "Methodology_Caveats"]
    wb._sheets = [wb[n] for n in order]; wb.active = 0

    os.makedirs(outdir, exist_ok=True)
    xlsx = os.path.join(outdir, f"Daewoo_x_ShootOrder-Deep_Analysis_Report-MTD-1-{end.day}{end.strftime('%b')}{end.year}.xlsx"); wb.save(xlsx)
    lo = os.path.join(outdir, '_lo'); subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', lo, xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    os.replace(os.path.join(lo, os.path.basename(xlsx)), xlsx); os.rmdir(lo)

    # ---------------- digest + facts
    fails = [n for n, ok, d in qa if not ok]
    lines = [f"# Daewoo × ShootOrder — Daily Digest | MTD 1–{label}", ""]
    if fails: lines += ["> ⚠️ **QA FAILED — do not forward until fixed:** " + "; ".join(fails), ""]
    lines += [f"**Funnel:** {inr(tot_sp)} spend → {A_leads:,} leads (CPL {inr(tot_sp/A_leads)}) → {A_q} qualified ({pc(A_q/A_leads)}, CPQL {inr(tot_sp/A_q)}) → {A_m} meetings → **{inr(cost_mtg)} per meeting**.", ""]
    if delta_rows:
        tdel = np.array([r_[1:] for r_ in delta_rows]).sum(axis=0)
        lines += [f"**Since last snapshot ({pd_}):** spend {inr(tdel[0])}, link clicks {int(tdel[2]):,}, net CRM leads {int(tdel[3])}, qualified {int(tdel[4])}, meetings {int(tdel[5])}.", ""]
    lines += ["## Key findings"] + [f"{i}. **{h}** {t}" for i, (h, t) in enumerate(F, 1)] + ["", "## Do next"] + [f"- {a[1]}" for a in acts[:6]]
    open(os.path.join(outdir, "digest.md"), "w").write("\n".join(lines) + "\n")
    facts = dict(date=stamp, spend=tot_sp, leads=A_leads, qualified=A_q, meetings=A_m, cost_per_meeting=cost_mtg, qa=[dict(check=n, ok=ok, detail=d) for n, ok, d in qa],
                 scale=sc_[['ad', 'adset', 'cpm']].to_dict('records'), reduce=rd_[['ad', 'adset', 'spend']].to_dict('records'))
    json.dump(facts, open(os.path.join(outdir, "facts.json"), "w"), indent=2, default=float)
    print(f"OK  {xlsx}\n    digest: {os.path.join(outdir,'digest.md')}\n    QA: {'ALL PASS' if not fails else 'FAIL -> ' + '; '.join(fails)}")
    return 0 if not fails else 2


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("report"); ap.add_argument("--outdir"); ap.add_argument("--history", default="history/daily_snapshot.csv")
    ap.add_argument("--targets", default="config/targets.json")
    a = ap.parse_args()
    out = a.outdir or os.path.join("outputs", os.path.basename(os.path.dirname(os.path.abspath(a.report))))
    sys.exit(build(a.report, out, a.history, a.targets))
