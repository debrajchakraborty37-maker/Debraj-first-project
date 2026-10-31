import copy
import pandas as pd
from datalayer import QUALITY
from build_helpers import clear_from, put
QTXT = 'Won, Token Received, Qualified, Meeting Scheduled, Meeting Conducted, Follow-Up, Details Shared'


def fill(wb, S, d, ct, act, f, B, k):
    df, L, E, X = d['df'], d['L'], d['E'], d['X']
    camp, adst, ads = d['camp'], d['adst'], d['ads']
    grp = df['Brand'].isin(d['cfg']['groups'])
    tag_ok = df['camp_res'].map(lambda x: isinstance(x, str))
    other = int((~grp & ~tag_ok).sum())
    n_lines = {'raw': d['n_raw'], 'dup': d['n_dup']}
    ws = S['Reconciliation']
    st = {c: copy.copy(ws.cell(5, c)._style) for c in range(1, 8)}
    hdr = {c: copy.copy(ws.cell(4, c)._style) for c in range(1, 8)}
    sec = copy.copy(ws['A26']._style); b1 = copy.copy(ws['B27']._style); b2 = copy.copy(ws['B28']._style); cv = copy.copy(ws['C27']._style)
    nt = copy.copy(ws['B51']._style); t1 = copy.copy(ws['A1']._style); t2 = copy.copy(ws['A2']._style)
    clear_from(ws, 5)
    ws['A2'] = 'Source figures are computed directly from the raw files provided (pandas, independent of the workbook); report figures are live formulas on the workbook tabs. Status = OK when difference is zero.'
    nL, nF, mc, nA, last = k['nL'], k['nF'], k['mc_last'], k['nA'], k['last']
    ex_n, xb_n = len(E), len(X)
    qual_all_group = int(df[grp & df['Status'].isin(QUALITY)].shape[0])
    qual_scope = int(L.qual.sum())
    blank_un = int((L.campaign == '(Unattributed)').sum())
    meta_res = float(ct[ct.in_meta].meta.sum())
    dsh_g = f"'Master Dashboard'!E{k['dash_end'] + 4}"
    rows = [
        ('CRM rows in LeadData_October2026.csv (all brands)  =  duplicate rows + in-scope + excluded + cross-brand + other-brand records', d['n_raw'],
         f"={n_lines['dup']}+COUNTA(Lead_Data!A2:A{nL + 1})+COUNTA(Excluded_Leads!A2:A{ex_n + 1})+COUNTA(XBrand_Leads!A2:A{xb_n + 1})+{other}",
         f"{d['n_raw']:,} rows: {d['n_dup']} duplicate Lead id rows removed + {nL:,} {B}-campaign leads + {ex_n} excluded {B}-brand rows + {xb_n} cross-brand + {other:,} other-brand records"),
        ('In-scope leads  =  Old + New + Unattributed', nL, '=COUNTIF(ldType,"Old")+COUNTIF(ldType,"New")+COUNTIF(ldType,"Unattributed")', 'Every in-scope lead is classified'),
        ('Fact_Data lead rows  =  Lead_Data rows', nL, '=COUNTIF(fRT,"Lead")', 'Pivot source table vs lead list'),
        ('Fact_Data quality leads  =  Lead_Data quality leads', qual_scope, '=COUNTIF(ldQual,"Quality Lead")', 'Quality = status in the Settings list'),
        (f'Quality-status records among all {B}-brand rows in the CRM file  =  in-scope + excluded', qual_all_group,
         f'=COUNTIF(ldQual,"Quality Lead")+COUNTIF(Excluded_Leads!H2:H{ex_n + 1},"Yes")', 'Every quality record of the brand is accounted for'),
        ('Leads with no campaign tag kept as Unattributed', blank_un, '=COUNTIF(ldType,"Unattributed")', 'Paid / WhatsApp-API leads with campaign blank - shown as Unattributed, not forced into Old/New'),
        ('Bulk / manual (Excel) uploads with no campaign - excluded', int(E.reason.str.startswith('No campaign').sum()), f'=COUNTIF(Excluded_Leads!I2:I{ex_n + 1},"No campaign*")', 'Not campaign leads (lead source = Excel)'),
        ('Ad spend: campaign file total (INR)', round(float(camp['Amount spent (INR)'].sum()), 2), '=SUM(fSpend)', 'Fact_Data spend rows (one per campaign) sum to the Meta campaign file'),
        ('Ad spend: ad-set file total (INR)', round(float(adst['Amount spent (INR)'].sum()), 2), '=SUM(fSpend)', 'Campaign file = ad-set file = report'),
        ('Ad spend: sum of ad sets mapped to campaigns  vs  campaign file (Campaign_Master total difference)', 0, '=Campaign_Master!J%d' % (last + 1), 'Inferred ad-set to campaign mapping reproduces every campaign total exactly (spend, impressions, results)'),
        ('Impressions: campaign file total', int(camp['Impressions'].sum()), '=SUM(fImpr)', ''),
        ('Link clicks: campaign file total', float(camp['Link clicks'].fillna(0).sum()), '=SUM(fClicks)', 'Report uses campaign-file clicks (the ad-set export sums to %d more in places)' % int(adst['Link clicks'].fillna(0).sum() - camp['Link clicks'].fillna(0).sum())),
        ('Meta lead-type results (campaign file, excl. landing-page-view campaigns)', meta_res, '=SUM(fMeta)', 'Web campaigns report landing page views, not leads'),
        ('Ad-level file spend (INR)  vs  campaign file', round(float(camp['Amount spent (INR)'].sum()), 2), round(float(ads['Amount spent (INR)'].sum()), 2),
         'Ad-level export may miss deleted / archived ads and repeats ad names across ad sets - ad-level spend is NOT used; spend is reported at campaign level'),
        ('Dashboard grand-total leads  =  Lead_Data leads', nL, f'={dsh_g}', 'Dashboard ties to the lead list'),
        ('Leads created outside the reporting window (01-Oct to 31-Oct)', 0, '=COUNTIF(ldDate,"<"&WinStart)+COUNTIF(ldDate,">"&WinEnd)', 'Lead export covers exactly the window'),
        (f'{B}-campaign leads tagged to OTHER brands in CRM (listed, not counted)', xb_n, f'=COUNTA(XBrand_Leads!A2:A{xb_n + 1})', 'Reported separately on XBrand_Leads; their status belongs to another brand pipeline'),
    ]
    for i, (chk, src, rep, com) in enumerate(rows):
        r = 5 + i
        put(ws, (r, 1), i + 1, st[1]); put(ws, (r, 2), chk, st[2]); put(ws, (r, 3), src, st[3]); put(ws, (r, 4), rep, st[3])
        put(ws, (r, 5), f'=ROUND(D{r}-C{r},2)', st[3]); put(ws, (r, 6), f'=IF(ABS(E{r})<=0.005,"OK","SEE NOTE")', st[6]); put(ws, (r, 7), com, st[7])
        ws.row_dimensions[r].height = 30
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting = type(ws.conditional_formatting)()
    tw = openpyxl_tmpl_cf()
    info_row = 5 + [i for i, x_ in enumerate(rows) if x_[0].startswith('Ad-level')][0]
    last_r = 4 + len(rows)
    for cf in tw:
        for rule in cf.rules: ws.conditional_formatting.add(f'F5:F{info_row - 1} F{info_row + 1}:F{last_r}', rule)
    ws.cell(info_row, 6).value = 'INFO'
    r = 5 + len(rows) + 1
    put(ws, (r, 1), 'LEAD FLOW  -  from raw file to report', sec); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    flow = [(f'Rows in LeadData_October2026.csv (all brands)', d['n_raw'], b1), ('  less: duplicate Lead id rows (same lead exported twice)', d['n_dup'], b2),
            ("  less: other brands' records (not a %s brand and not on a %s campaign)" % (B, B), other, b2),
            (f'  less: {B}-brand rows excluded (bulk upload / other-brand campaign)', ex_n, b2),
            (f'  less: {B}-campaign leads tagged to other brands (XBrand_Leads)', xb_n, b2),
            (f'{B}-campaign leads used in the report (total)', nL, b1),
            ('     of which Old campaigns', int((L.ctype == 'Old').sum()), b2), ('     of which New campaigns', int((L.ctype == 'New').sum()), b2),
            ('     of which Unattributed', blank_un, b2),
            ('     memo: New-campaign leads from EBG - In House campaigns (no Meta spend data)', int(k['nospend_leads']), b2)]
    for j, (lab, v, s_) in enumerate(flow):
        put(ws, (r + 1 + j, 2), lab, s_); put(ws, (r + 1 + j, 3), v, cv)
    r2 = r + len(flow) + 2
    put(ws, (r2, 2), ('No separate Lead Center export was supplied for this period; LeadData_October2026.csv (CRM export, one row per Lead id) is the single lead source. '
                      'The previous reports\' Lead Center reconciliation and Performance-Report cross-checks are therefore not repeated.'), nt)
    ws.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=7); ws.row_dimensions[r2].height = 32

    # ------------------------------------------------------------ Notes
    ws = S['Notes']
    sty_h = copy.copy(ws['B3']._style); sty_b = copy.copy(ws['B4']._style); sty_t = copy.copy(ws['B1']._style)
    clear_from(ws, 3)
    O, N = k['O'], k['N']
    alias_n = int((L.method.str.startswith('Alias')).sum())
    lines = [
        ('h', 'HOW TO READ THIS WORKBOOK'),
        ('b', f'•  Master Dashboard: the one-page management view (formula-driven): Old vs New headline KPIs, quality parameters, and a campaign-wise table with type, spend, leads, quality count, quality %, CPL and CPQL. '
              'Master Pivot: real PivotTables + slicers (Campaign Type, Campaign, Lead Status, Month, Week) plus the Campaign x Quality Parameter table at the bottom. Meta spend cannot be split by status / month / week, so spend shows as "-" when those slicers are used.'),
        ('b', '•  Executive Summary: takeaways, Old-vs-New KPI table, status funnel, same-age cohort and campaign scorecard - all live formulas.'),
        ('b', '•  Reconciliation: every figure tied back to the raw files. Campaign_Master: Old/New classification (editable). Data tabs: Fact_Data (pivot source), Lead_Data, Excluded_Leads, XBrand_Leads, Meta_* (raw Meta exports), Settings.'),
        ('h', 'DEFINITIONS'),
        ('b', f'•  Quality Lead = lead whose CRM Status is one of: {QTXT} (editable on Settings). "Meeting Conducted but NOT QUALIFIED" is not a quality status.'),
        ('b', "•  Old / New: a campaign is NEW if its name carries 'Oct26' - the new October-2026 campaign structure (ABO campaigns in the Kustard ad account and the 'EBG - In House ... Oct26' campaigns in the In-House ad account). Every other campaign is Old. Editable on Campaign_Master."),
        ('b', '•  CRM leads = lead records in the CRM (one row per Lead id). Contacted = status other than "New Lead". Quality % = Quality leads / CRM leads. CPL = ad spend / CRM leads. CPQL = ad spend / Quality leads.'),
        ('b', '•  Meta results = instant-form leads (leadgen) plus WhatsApp conversations started. Web landing-page campaigns report landing page views (not leads) - excluded from results but their spend is included in spend.'),
        ('h', 'DATA TREATMENT & FINDINGS'),
        ('b', f"•  REPORTING WINDOW: the Meta exports (campaigns, ad sets, ads) and LeadData_October2026.csv cover 01-Oct-2026 to 31-Oct-2026 only (the CRM export was taken on 31-Oct, last lead created 31-Oct 12:56). "
              "September is not in the supplied files, so this report compares Old vs New on October data; it is not cumulative with the earlier 01-Sep to 18-Oct report."),
        ('b', f"•  LeadData_October2026.csv holds {d['n_raw']:,} rows for all brands ({d['n_dup']} duplicate Lead id rows removed). {nL:,} leads belong to {B} campaigns and are analysed; {other:,} records belong to other brands and are out of scope; "
              f"{ex_n} {B}-brand rows are excluded (see Excluded_Leads) and {xb_n} {B}-campaign leads tagged to other brands are listed on XBrand_Leads but not counted."),
        ('b', f"•  SPEND COVERAGE (important): the EBG - In House campaigns ({', '.join(sorted(set(k['nospend'])))[:230]}) have {k['nospend_leads']:,} CRM leads but are NOT in the Meta exports supplied (they run in a different ad account), so their spend is unknown. "
              "Their leads and quality are counted in lead / quality metrics, but cost metrics (CPL, CPQL, capture rate) are computed only on leads of campaigns that have Meta spend data (Fact_Data columns Costed Lead Count / Costed Quality Count). "
              "Send the EBG - In House Meta exports to add their spend; no other change is needed."),
        ('b', f"•  Ad-set names are reused across campaigns and the ad-set export has no campaign column. Ad sets were mapped to campaigns by exact-sum matching (spend, impressions and results of every campaign are reproduced exactly) with CRM ad-set evidence as tie-breaker. "
              "Because identical ad-set names exist in several campaigns, an individual ad-set's campaign is an inference; all report figures use CAMPAIGN-level Meta data and are not affected."),
        ('b', f"•  The ad-level Meta export totals INR {ads['Amount spent (INR)'].sum():,.2f} vs INR {camp['Amount spent (INR)'].sum():,.2f} in the campaign export, and ad names repeat across ad sets - so spend is reported at campaign level only; ad-level analysis uses lead counts and quality."),
    ]
    if alias_n:
        lines.append(('b', f"•  {alias_n} leads carry the CRM campaign name 'ABO | CW | WhatsApp | Unit Franchise | India | Oct26 | V2', which does not exist in Meta (the only CW WhatsApp campaign there is '... Oct26 | V1'); they were mapped to that V1 campaign (Attribution Method = Alias)."))
    inf = int(L.method.str.startswith('Inferred').sum())
    if inf:
        lines.append(('b', f"•  {inf} WhatsApp-API leads have a blank campaign in the CRM but are tagged to the brand 'In House Carlton Salon' (the EBG - In House account). As instructed, they are classified as New under the pseudo-campaign '{L[L.method.str.startswith('Inferred')].campaign.iloc[0]}' (Attribution Method = Inferred); this is a brand-based assumption, not a CRM campaign tag, and their spend is not available."))
    if blank_un:
        lines.append(('b', f"•  {blank_un} paid / WhatsApp-API leads have a blank campaign in the CRM and no ad-set information. They are shown as Unattributed (not forced into Old or New); {int(L[(L.campaign=='(Unattributed)') & (L.status=='New Lead')].shape[0])} of them are still 'New Lead'."))
    lines += [
        ('b', f"•  {int(E.reason.str.startswith('No campaign').sum())} {B}-brand rows came from bulk / manual (Excel) uploads with no campaign or ad set and are excluded; they are listed on Excluded_Leads."),
        ('b', '•  CRM attribution is taken exactly as tagged (campaign tag = source of truth).'),
        ('h', 'IMPORTANT CAVEATS'),
        ('b', f"•  Lead ageing: New-campaign leads are at most {k['new_days']} days old at the 31-Oct snapshot, so many are still in early contact. Compare on 'Quality % of contacted' and the same-age cohort (Executive Summary section 3) for a fair view. New-campaign quality will keep maturing."),
        ('b', '•  Old campaigns carry different budgets / audiences / launch dates; New campaigns are a single audience-test structure (Investors, Business Owners, Broad, LAL 1%). Totals are not scaled; daily averages are provided as a normalising view.'),
        ('b', '•  Phone numbers are masked and names / e-mails are not included in this workbook.'),
    ]
    put(ws, (1, 2), f'Notes, Definitions & Assumptions - {B}', sty_t)
    r = 3
    for kind, text in lines:
        put(ws, (r, 2), text, sty_h if kind == 'h' else sty_b)
        ws.row_dimensions[r].height = 15.5 if kind == 'h' else max(18, 15 * (len(text) // 135 + 1))
        r += 1


def openpyxl_tmpl_cf():
    import openpyxl, warnings
    warnings.filterwarnings('ignore')
    TEMPLATE='/root/.claude/uploads/eb8fa0cd-ad52-5e00-bd26-b0914f0c0cda/00143e64-Natuf_Old_vs_New_Campaign_Lead_Quality_Report.xlsx'
    return openpyxl.load_workbook(TEMPLATE)['Reconciliation'].conditional_formatting
